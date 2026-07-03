from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ABAS_PESQUISA = ["P1", "P2", "P3", "P4"]
FORMATO_DATA = "dd/mm/yyyy"
ULTIMA_LINHA_FORMULA_PESQUISA = 12000
ULTIMA_LINHA_FORMULA_STATUS = 70000


FORMULAS_BASE = [
    {
        "coluna": "DT ENVIO",
        "letra": "AB",
        "formula": lambda linha: (
            f'=SEERRO(SE(AI{linha}<>"",AI{linha},'
            f'PROCX(F{linha},STATUS!U:U,STATUS!R:R,"")),"")'
        ),
    },
    {
        "coluna": "CHAVE",
        "letra": "AF",
        "formula": lambda linha: f'=J{linha}&"_"&V{linha}&"_"&Y{linha}&"_"&AA{linha}&"_"&F{linha}',
    },
    {
        "coluna": "STATUS",
        "letra": "AM",
        "formula": lambda linha: (
            f'=SE(OU(AK{linha}="Não quis",AK{linha}="Óbito",AK{linha}="Lida"),'
            f'AK{linha},SE(CONT.SES(STATUS!U:U,F{linha},STATUS!G:G,"Lida")>0,"Lida",""))'
        ),
    },
    {
        "coluna": "P1",
        "letra": "AO",
        "formula": lambda linha: f'=PROCX(F{linha},\'P1\'!R:R,\'P1\'!E:E,"")',
    },
    {
        "coluna": "P2",
        "letra": "AP",
        "formula": lambda linha: f'=PROCX(F{linha},\'P2\'!R:R,\'P2\'!E:E,"")',
    },
    {
        "coluna": "P3",
        "letra": "AQ",
        "formula": lambda linha: f'=PROCX(F{linha},\'P3\'!R:R,\'P3\'!E:E,"")',
    },
    {
        "coluna": "P4",
        "letra": "AR",
        "formula": lambda linha: f'=PROCX(F{linha},\'P4\'!R:R,\'P4\'!E:E,"")',
    },
]


def ultima_linha_com_dados(ws, coluna_referencia="J"):
    for linha in range(ws.max_row, 1, -1):
        valor = ws[f"{coluna_referencia}{linha}"].value
        if valor not in (None, ""):
            return linha
    return 1


def validar_posicoes_colunas(ws):
    erros = []

    for regra in FORMULAS_BASE:
        valor_cabecalho = ws[f"{regra['letra']}1"].value

        if valor_cabecalho == regra["coluna"]:
            print(f"OK formula: {regra['coluna']} encontrada em {regra['letra']}")
        else:
            mensagem = (
                f"ERRO formula: esperado {regra['coluna']} em {regra['letra']}, "
                f"mas encontrado {valor_cabecalho}"
            )
            print(mensagem)
            erros.append(mensagem)

    if erros:
        raise ValueError("Posicionamento de colunas invalido para aplicar formulas.")


def localizar_coluna_por_cabecalho(ws, cabecalho):
    for celula in ws[1]:
        if celula.value == cabecalho:
            return celula.column
    raise ValueError(f"Coluna {cabecalho} nao encontrada na aba {ws.title}.")


def formula_apos_ultimo_underscore(celula):
    return (
        f'=IFERROR(RIGHT({celula},LEN({celula})-'
        f'FIND("@",SUBSTITUTE({celula},"_","@",'
        f'LEN({celula})-LEN(SUBSTITUTE({celula},"_",""))))),"")'
    )


def aplicar_formulas_senha_pesquisa(wb):
    for nome_aba in ABAS_PESQUISA:
        ws = wb[nome_aba]
        coluna_nome = localizar_coluna_por_cabecalho(ws, "Nome")
        coluna_senha = localizar_coluna_por_cabecalho(ws, "SENHA")
        letra_nome = get_column_letter(coluna_nome)
        letra_senha = get_column_letter(coluna_senha)

        for linha in range(2, ULTIMA_LINHA_FORMULA_PESQUISA + 1):
            ws[f"{letra_senha}{linha}"] = formula_apos_ultimo_underscore(f"{letra_nome}{linha}")

        print(
            f"Formula aplicada em {nome_aba}!{letra_senha}2:"
            f"{letra_senha}{ULTIMA_LINHA_FORMULA_PESQUISA} para extrair SENHA de Nome"
        )


def aplicar_formula_senha_status(wb):
    ws = wb["STATUS"]
    coluna_contato = localizar_coluna_por_cabecalho(ws, "Contato")
    coluna_senha = localizar_coluna_por_cabecalho(ws, "SENHA")
    letra_contato = get_column_letter(coluna_contato)
    letra_senha = get_column_letter(coluna_senha)

    for linha in range(2, ULTIMA_LINHA_FORMULA_STATUS + 1):
        ws[f"{letra_senha}{linha}"] = formula_apos_ultimo_underscore(f"{letra_contato}{linha}")

    print(
        f"Formula aplicada em STATUS!{letra_senha}2:"
        f"{letra_senha}{ULTIMA_LINHA_FORMULA_STATUS} para extrair SENHA de Contato"
    )


def aplicar_formatos_base(ws):
    coluna_data_internacao = localizar_coluna_por_cabecalho(ws, "DT INTERNACAO")
    letra_data_internacao = get_column_letter(coluna_data_internacao)

    for linha in range(2, ws.max_row + 1):
        ws[f"{letra_data_internacao}{linha}"].number_format = FORMATO_DATA

    print(
        f"Formato aplicado em DT INTERNACAO "
        f"({letra_data_internacao}2:{letra_data_internacao}{ws.max_row})"
    )


def aplicar_formulas(caminho_arquivo):
    wb = load_workbook(caminho_arquivo)
    ws = wb["BASE"]

    validar_posicoes_colunas(ws)
    ultima_linha = ultima_linha_com_dados(ws)

    if ultima_linha < 2:
        print("Nenhuma linha de dados encontrada para aplicar formulas.")
        wb.save(caminho_arquivo)
        return

    for regra in FORMULAS_BASE:
        for linha in range(2, ultima_linha + 1):
            ws[f"{regra['letra']}{linha}"] = regra["formula"](linha)

        print(f"Formula aplicada em {regra['coluna']} ({regra['letra']}2:{regra['letra']}{ultima_linha})")

    aplicar_formulas_senha_pesquisa(wb)
    aplicar_formula_senha_status(wb)
    aplicar_formatos_base(ws)

    wb.save(caminho_arquivo)
